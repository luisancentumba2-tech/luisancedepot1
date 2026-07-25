"""Construit Morning_SARL_Systeme_Paie_RH_RDC.xlsm (version macro) à partir du
classeur .xlsx (version sans macro, canonique).

Approche : le projet est construit sur Linux, sans Excel/COM disponible, donc
impossible d'utiliser l'éditeur VBA pour compiler un vbaProject.bin. Ce script
construit le fichier OLE2/MS-CFB vbaProject.bin directement en Python :

  - vba/vbaProject_template.bin : vbaProject.bin réel, généré par Excel, tiré
    des exemples officiels de la bibliothèque XlsxWriter (BSD, prévu pour cet
    usage : voir examples/vba_extract.py dans XlsxWriter). Fournit une
    structure de classeur VBA déjà valide (PROJECT, PROJECTwm, dir, modules
    document ThisWorkbook/Sheet1/Sheet2/ThisWorkbook1, module standard
    Module1). Ces modules "document" ne correspondent à aucune feuille réelle
    de ce classeur (8 feuilles, pas 2) — c'est sans risque : XlsxWriter
    distribue ce même template pour être attaché à des classeurs de n'importe
    quelle taille.
  - vba/ovba_compress.py : implémente la compression MS-OVBA (2.4.1) en mode
    "littéral seul" (pas de rétro-référence) — suffisant pour la correction,
    pas pour la taille. Vérifié par aller-retour via le décompresseur de
    référence d'oletools (voir son bloc `__main__`).
  - vba/cfb_writer.py : reconstruit le conteneur OLE2 en repartant de zéro
    (FAT, MiniFAT, mini-flux, arbre de répertoire) plutôt que de patcher le
    template en place, car `olefile` ne sait réécrire un flux qu'à taille
    identique.
  - vba/ClotureDuMois.bas : code source de la macro, seule source de vérité
    (modifier ce fichier, pas le binaire).

Le module "Module1" du template est remplacé par le code de ClotureDuMois.bas
avec TextOffset=0 (flux = uniquement le code source compressé, sans P-code
préalable) — ceci pour éviter le motif de "VBA stomping" (P-code et source
qui ne correspondent pas) que produirait la conservation du P-code d'origine
du template, motif que certains antivirus signalent à tort comme suspect.

Vérification effectuée (voir conversation de session) : extraction du code
via `oletools.olevba` sur le .bin ET sur le .xlsm final → correspond
octet pour octet au contenu de ClotureDuMois.bas. Structure OLE2 validée par
`olefile`. Chargement complet du .xlsm validé par `openpyxl` (mêmes noms
définis, mêmes formules, mêmes validations de données que le .xlsx source).
AUCUN test n'a pu être fait dans Excel réel (environnement Linux sans COM) —
à faire avant tout usage en production.

Dépendances (pip) : openpyxl, olefile, oletools (ce dernier seulement pour la
vérification, pas requis pour la construction elle-même).
"""
import re
import shutil
import subprocess
import sys
import zipfile
from copy import copy
from pathlib import Path

import openpyxl
import olefile

sys.path.insert(0, str(Path(__file__).parent / "vba"))
from ovba_compress import compress_stream  # noqa: E402

HERE = Path(__file__).parent
REPO = HERE.parent
SOURCE_XLSX = REPO / "Morning_SARL_Systeme_Paie_RH_RDC.xlsx"
OUTPUT_XLSM = REPO / "Morning_SARL_Systeme_Paie_RH_RDC.xlsm"
TEMPLATE_BIN = HERE / "vba" / "vbaProject_template.bin"
MACRO_SOURCE = HERE / "vba" / "ClotureDuMois.bas"

INSTRUCTION_NOTE = (
    "⚙  Version .xlsm de ce classeur : la clôture ci-dessus peut être automatisée avec la macro "
    "ClotureDuMois (Alt+F8 → ClotureDuMois → Exécuter). Elle applique exactement les étapes 1 à 5 "
    "ci-dessus (copie figée de JOURNAL, statut Clos, date, puis avance le mois de PARAMÈTRES) et "
    "refuse de ré-écraser un bloc déjà clos. Macro non testée dans Excel réel (construite hors "
    "environnement Windows) — à vérifier avant usage en production."
)


def add_instruction_note(src_xlsx: Path, dst_xlsx: Path) -> None:
    wb = openpyxl.load_workbook(src_xlsx, data_only=False)
    ws = wb["HISTORIQUE"]
    if "B11:R11" not in [str(r) for r in ws.merged_cells.ranges]:
        ws.merge_cells("B11:R11")
    b10, b11 = ws["B10"], ws["B11"]
    b11.value = INSTRUCTION_NOTE
    b11.font = copy(b10.font)
    b11.alignment = copy(b10.alignment)
    ws.row_dimensions[11].height = 40
    wb.save(dst_xlsx)


def build_vba_project() -> bytes:
    from cfb_writer import build_compound_file
    import oletools.olevba as olevba
    import struct

    ole = olefile.OleFileIO(str(TEMPLATE_BIN))

    # Patch the `dir` stream: set Module1's MODULEOFFSET (TextOffset) to 0.
    dir_plain = bytearray(olevba.decompress_stream(bytearray(ole.openstream("VBA/dir").read())))
    pattern = struct.pack("<HI", 0x0031, 4) + struct.pack("<I", 955)
    hits = [i for i in range(len(dir_plain)) if dir_plain[i:i + len(pattern)] == pattern]
    if len(hits) != 1:
        raise RuntimeError(f"expected exactly one Module1 MODULEOFFSET(955) record, found {len(hits)}")
    value_offset = hits[0] + 6
    dir_plain[value_offset:value_offset + 4] = struct.pack("<I", 0)
    new_dir = compress_stream(bytes(dir_plain))

    src_text = MACRO_SOURCE.read_text(encoding="utf-8").replace("\r\n", "\n").replace("\n", "\r\n")
    new_module1 = compress_stream(src_text.encode("cp1252"))

    streams = [
        {"name": "PROJECT", "storage": "root", "data": ole.openstream("PROJECT").read()},
        {"name": "PROJECTwm", "storage": "root", "data": ole.openstream("PROJECTwm").read()},
    ]
    for nm in ["ThisWorkbook", "Sheet1", "Sheet2", "ThisWorkbook1",
               "_VBA_PROJECT", "__SRP_0", "__SRP_1", "__SRP_2", "__SRP_3"]:
        streams.append({"name": nm, "storage": "VBA", "data": ole.openstream("VBA/" + nm).read()})
    streams.append({"name": "dir", "storage": "VBA", "data": new_dir})
    streams.append({"name": "Module1", "storage": "VBA", "data": new_module1})

    return build_compound_file(streams)


def inject_vba_project(xlsx_path: Path, vba_bytes: bytes, xlsm_path: Path) -> None:
    build_dir = HERE / "_xlsm_build"
    if build_dir.exists():
        shutil.rmtree(build_dir)
    build_dir.mkdir()
    with zipfile.ZipFile(xlsx_path) as z:
        z.extractall(build_dir)

    (build_dir / "xl" / "vbaProject.bin").write_bytes(vba_bytes)

    ct_path = build_dir / "[Content_Types].xml"
    ct = ct_path.read_text(encoding="utf-8")
    ct = ct.replace(
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"',
        'ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml"',
    )
    ct = ct.replace(
        "</Types>",
        '<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
    )
    ct_path.write_text(ct, encoding="utf-8")

    rels_path = build_dir / "xl" / "_rels" / "workbook.xml.rels"
    rels = rels_path.read_text(encoding="utf-8")
    used_ids = [int(m) for m in re.findall(r'Id="rId(\d+)"', rels)]
    new_id = max(used_ids) + 1 if used_ids else 1
    rels = rels.replace(
        "</Relationships>",
        f'<Relationship Id="rId{new_id}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vbaProject" '
        'Target="vbaProject.bin"/></Relationships>',
    )
    rels_path.write_text(rels, encoding="utf-8")

    if xlsm_path.exists():
        xlsm_path.unlink()
    with zipfile.ZipFile(xlsm_path, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(build_dir.rglob("*")):
            if f.is_file():
                z.write(f, f.relative_to(build_dir))

    shutil.rmtree(build_dir)


def verify(xlsm_path: Path) -> None:
    import oletools.olevba as olevba

    wb = openpyxl.load_workbook(xlsm_path, keep_vba=True)
    assert wb.vba_archive is not None, "no VBA archive found after load"

    with zipfile.ZipFile(xlsm_path) as z:
        assert z.testzip() is None
        vba_bytes = z.read("xl/vbaProject.bin")

    ole = olefile.OleFileIO(vba_bytes)
    proj = olevba.VBA_Project(ole, "", "PROJECT", "VBA/dir", relaxed=True)
    list(proj.parse_modules())
    module1 = next(m for m in proj.modules if m.name == "Module1")
    expected = MACRO_SOURCE.read_text(encoding="utf-8").replace("\r\n", "\n")
    extracted = module1.code.replace("\r\n", "\n")
    if extracted != expected:
        raise RuntimeError("extracted macro source does not match ClotureDuMois.bas")
    print("OK — vba_archive present, zip intègre, code macro extrait identique à ClotureDuMois.bas")


def main():
    tmp_xlsx = HERE / "_tmp_with_note.xlsx"
    add_instruction_note(SOURCE_XLSX, tmp_xlsx)
    vba_bytes = build_vba_project()
    inject_vba_project(tmp_xlsx, vba_bytes, OUTPUT_XLSM)
    tmp_xlsx.unlink()
    verify(OUTPUT_XLSM)
    print(f"écrit {OUTPUT_XLSM}")


if __name__ == "__main__":
    main()
